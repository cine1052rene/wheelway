"""WheelWay - 실기기 스크린샷 브리핑 PPT 생성 스크립트"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

GREEN = RGBColor(0x08, 0x70, 0x5B)      # WheelWay 브랜드 메인 컬러
ORANGE = RGBColor(0xB4, 0x54, 0x00)     # 브랜드 보조 컬러
NAVY = RGBColor(0x1B, 0x2A, 0x2A)
GRAY = RGBColor(0x55, 0x5B, 0x66)
LIGHT = RGBColor(0xF1, 0xF2, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

IMG_DIR = r"C:\Users\cine1\AppData\Local\Temp\wheelway-screens"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def add_bg(slide, color=WHITE):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    slide.shapes._spTree.remove(shp._element)
    slide.shapes._spTree.insert(2, shp._element)
    return shp


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=NAVY,
                 align=PP_ALIGN.LEFT, font="맑은 고딕", anchor=None, line_spacing=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return tb


def add_footer(slide, page_no):
    add_textbox(slide, 0.5, 7.08, 8, 0.35, "WheelWay · 앱 화면 브리핑", size=10, color=GRAY)
    add_textbox(slide, 12.0, 7.08, 0.8, 0.35, str(page_no), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def bar(slide, left, top, width, height, color=GREEN):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def rounded_box(slide, left, top, width, height, fill=LIGHT, line_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def phone_shot(slide, filename, left, top, height=6.2):
    """실기기 세로 스크린샷(1440x3088)을 비율 유지해서 배치하고 얇은 테두리를 준다."""
    width = height * (1440 / 3088)
    border = rounded_box(slide, left - 0.06, top - 0.06, width + 0.12, height + 0.12,
                          fill=WHITE, line_color=RGBColor(0xD8, 0xDC, 0xD8))
    pic = slide.shapes.add_picture(f"{IMG_DIR}\\{filename}", Inches(left), Inches(top), height=Inches(height))
    return pic


def content_slide(page_no, kicker, title, desc_lines, filename, tag=None, tag_color=GREEN):
    slide = add_slide()
    add_bg(slide)
    bar(slide, 0, 0, 13.333, 0.12, GREEN)

    add_textbox(slide, 0.7, 0.5, 6.8, 0.4, kicker, size=14, bold=True, color=ORANGE)
    add_textbox(slide, 0.7, 0.9, 6.8, 0.9, title, size=30, bold=True, color=NAVY)

    if tag:
        tagbox = rounded_box(slide, 0.7, 1.75, 2.6, 0.5, fill=tag_color)
        tf = tagbox.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = tag
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.name = "맑은 고딕"
        desc_top = 2.5
    else:
        desc_top = 1.9

    desc_box = rounded_box(slide, 0.7, desc_top, 6.9, 4.3, fill=LIGHT)
    tb = slide.shapes.add_textbox(Inches(1.0), Inches(desc_top + 0.35), Inches(6.3), Inches(3.6))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(desc_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {line}"
        p.line_spacing = 1.3
        p.space_after = Pt(10)
        for r in p.runs:
            r.font.size = Pt(16)
            r.font.color.rgb = NAVY
            r.font.name = "맑은 고딕"

    phone_shot(slide, filename, left=8.5, top=0.75, height=6.2)
    add_footer(slide, page_no)
    return slide


# ── 표지 ──────────────────────────────────────────────
slide = add_slide()
add_bg(slide, GREEN)
add_textbox(slide, 1.0, 2.7, 11.3, 1.0, "WheelWay", size=54, bold=True, color=WHITE)
add_textbox(slide, 1.0, 3.6, 11.3, 0.7, "교통약자를 위한 지하철 접근성 안내 앱 — 화면 브리핑",
            size=22, color=WHITE)
add_textbox(slide, 1.0, 6.6, 11.3, 0.5, "실기기(갤럭시) 검증 스크린샷 기준", size=14,
            color=RGBColor(0xDD, 0xEE, 0xE6))

# ── 1. 홈 화면 ──────────────────────────────────────────
content_slide(
    2, "01  HOME", "지름길 찾기 — 첫 화면",
    [
        "앱을 열면 바로 뜨는 기본 화면으로, 출발역과 도착역만 고르면 된다.",
        "보행자 / 수동 휠체어 / 전동 휠체어 3가지 이동 수단 모드를 상단에서 바로 전환할 수 있다.",
        "복잡한 메뉴 없이 첫 화면에서 바로 목적을 시작하도록 설계했다.",
    ],
    "1_home.png",
    tag="이동수단 3종 지원", tag_color=GREEN,
)

# ── 2. 노선도 - 입력 전 ─────────────────────────────────
content_slide(
    3, "02  노선도 · 입력 전", "노선도에서 역을 고르기 전",
    [
        "네이버지도 스타일의 세로형 노선도로, 호선별 색상과 환승역 표시를 한눈에 볼 수 있다.",
        "역명을 몰라도 노선을 따라 훑어보며 원하는 역을 찾을 수 있다.",
        "상단 호선 칩을 누르면 해당 구간으로 빠르게 이동한다.",
    ],
    "2_linemap_before.png",
    tag="역 선택 전", tag_color=ORANGE,
)

# ── 3. 노선도 - 입력 후 ─────────────────────────────────
content_slide(
    4, "03  노선도 · 입력 후", "역을 탭하면 뜨는 선택 창",
    [
        "노선도에서 역 이름을 탭하면 하단에 출발역/도착역 지정 창이 바로 뜬다.",
        "화면 전환 없이 그 자리에서 바로 경로 검색에 반영된다.",
        "출발/도착을 고르는 즉시 첫 화면으로 자동 이동해 결과를 준비한다.",
    ],
    "3_linemap_after.png",
    tag="역 선택 직후", tag_color=GREEN,
)

# ── 4. 경로 결과 ────────────────────────────────────────
content_slide(
    5, "04  경로 검색 결과", "서울역 → 강남, 지름길 안내",
    [
        "소요시간, 환승 횟수, 구간별 노선 색상을 상단에 요약해서 보여준다.",
        "역 진입부터 환승, 도착까지 각 단계를 걸음 단위로 안내한다.",
        "구간마다 엘리베이터 위치 정보 유무까지 함께 표시해 교통약자 이동 계획에 도움을 준다.",
    ],
    "14_route_result.png",
    tag="핵심 기능", tag_color=GREEN,
)

# ── 5. 편의시설 조회 입력 ───────────────────────────────
content_slide(
    6, "05  편의시설 현황", "역명으로 편의시설 조회하기",
    [
        "역 이름만 입력하면 해당 역의 엘리베이터·에스컬레이터 위치를 조회할 수 있다.",
        "엘리베이터 / 에스컬레이터 필터로 원하는 시설만 골라볼 수 있다.",
        "공공데이터를 서버에서 미리 정리해두고 앱은 결과만 빠르게 받아온다.",
    ],
    "11_station_access.png",
    tag="입력 화면", tag_color=ORANGE,
)

# ── 6. 편의시설 조회 결과 ───────────────────────────────
content_slide(
    7, "06  편의시설 현황 결과", "강남역 엘리베이터 위치 안내",
    [
        "출구 번호, 정원(하중), 인접 출구 방향까지 실제 데이터로 안내한다.",
        "최근 서버 응답 속도를 크게 개선해 결과가 즉시 나타난다.",
        "실시간 도착정보 등 추가 정보도 같은 화면에서 함께 제공할 예정이다.",
    ],
    "18_facility_after_fix.png",
    tag="실데이터 확인", tag_color=GREEN,
)

# ── 마무리 ──────────────────────────────────────────────
slide = add_slide()
add_bg(slide, RGBColor(0xF1, 0xF2, 0xF1))
bar(slide, 0, 0, 13.333, 0.12, GREEN)
add_textbox(slide, 0.9, 1.2, 11.5, 0.8, "정리", size=32, bold=True, color=NAVY)
add_textbox(
    slide, 0.9, 2.2, 11.5, 4.0,
    "•  실기기(갤럭시)에서 직접 검증한 화면으로 구성했다.\n"
    "•  홈 → 노선도 → 경로 결과 → 편의시설 조회까지 핵심 사용 흐름을 담았다.\n"
    "•  화면을 확인하는 과정에서 실기기 전용 네트워크 문제를 발견해 즉시 수정했다.\n"
    "•  다음 단계: 실시간 도착정보 연동, 스토어 배포 준비.",
    size=18, color=NAVY, line_spacing=1.5,
)
add_footer(slide, 8)

OUT = r"C:\project\Wheelway\docs\reports\WheelWay_Screenshot_Briefing.pptx"
prs.save(OUT)
print("saved:", OUT)
